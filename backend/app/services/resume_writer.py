import io

from docx import Document

from app.models.schemas import Suggestion


def apply_suggestions_to_docx(
    file_bytes: bytes, suggestions: list[Suggestion], selected_ids: list[str]
) -> bytes:
    selected = {s.id for s in suggestions if s.id in selected_ids}
    replacements = [
        (s.original_text, s.suggested_text)
        for s in suggestions
        if s.id in selected
    ]

    if not replacements:
        return file_bytes

    doc = Document(io.BytesIO(file_bytes))

    for paragraph in doc.paragraphs:
        for original, suggested in replacements:
            if original in paragraph.text:
                paragraph.text = paragraph.text.replace(original, suggested, 1)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for original, suggested in replacements:
                        if original in paragraph.text:
                            paragraph.text = paragraph.text.replace(
                                original, suggested, 1
                            )

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
