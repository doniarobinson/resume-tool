import io

from docx import Document

from app.models.schemas import ResumeBlock, ResumeDocument

HEADING_STYLES = {
    "Title",
    "Heading 1",
    "Heading 2",
    "Heading 3",
    "Heading 4",
}


def _is_heading(paragraph) -> bool:
    style_name = paragraph.style.name if paragraph.style else ""
    if style_name in HEADING_STYLES:
        return True
    text = paragraph.text.strip()
    if not text:
        return False
    if len(text) < 80 and paragraph.runs:
        bold_runs = sum(1 for r in paragraph.runs if r.bold)
        if bold_runs >= max(1, len(paragraph.runs) // 2):
            return True
    return False


def parse_docx(file_bytes: bytes) -> ResumeDocument:
    doc = Document(io.BytesIO(file_bytes))
    blocks: list[ResumeBlock] = []
    current_section = "General"
    block_index = 0

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        if _is_heading(paragraph):
            current_section = text
            continue

        blocks.append(
            ResumeBlock(
                section=current_section,
                text=text,
                block_index=block_index,
            )
        )
        block_index += 1

    if not blocks:
        joined = "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
        if joined:
            blocks.append(ResumeBlock(section="General", text=joined, block_index=0))

    full_text = "\n\n".join(b.text for b in blocks)
    return ResumeDocument(full_text=full_text, blocks=blocks)
